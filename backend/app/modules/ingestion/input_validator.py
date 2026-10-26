"""Validation trên bản capture: một lần parse, không cleaning hay đổi giá trị nghiệp vụ."""

from __future__ import annotations

import csv
import sys
import time
from itertools import islice
from typing import TYPE_CHECKING, Any, Literal
from uuid import uuid4
from zipfile import BadZipFile

import psycopg
import pyarrow as pa
import pyarrow.parquet as pq
from minio.error import S3Error
from openpyxl import load_workbook
from psycopg import sql

from app.modules.ingestion.validation import (
    CSVOptions,
    ErrorCode,
    ExcelOptions,
    PostgreSQLOptions,
    RuleCheck,
    ValidationIssue,
    ValidationPolicy,
    ValidationReport,
)

if TYPE_CHECKING:
    from pathlib import Path

    from app.modules.ingestion.readers.contract import ReadRequest


class InputValidationError(Exception):
    """Mã lỗi đã kiểm soát; không chứa exception/giá trị thô của connector."""

    def __init__(self, rule: str, code: ErrorCode, *, infrastructure: bool = False) -> None:
        super().__init__(code)
        self.rule, self.code, self.infrastructure = rule, code, infrastructure


class ValidationBlockedError(Exception):
    def __init__(self, report: ValidationReport) -> None:
        super().__init__("BRONZE_INPUT_VALIDATION_BLOCKED")
        self.report = report


def required_rule_ids(source_type: str) -> tuple[str, ...]:
    prefix, count = {"CSV": ("CSV", 5), "EXCEL": ("XLS", 7), "POSTGRESQL": ("PG", 6)}[source_type]
    return (
        *(f"IN-COM-{i:03}" for i in range(1, 9)),
        *(f"IN-{prefix}-{i:03}" for i in range(1, count + 1)),
    )


class InputValidator:
    """Giới hạn theo scan/buffer, không phải sandbox giới hạn RSS toàn tiến trình."""

    def __init__(self, policy: ValidationPolicy, job_id: int) -> None:
        self.policy, self.job_id = policy, job_id
        self.started = time.monotonic()
        self.checks = {
            rule: RuleCheck(rule_id=rule, status="NOT_RUN")
            for rule in required_rule_ids(policy.parser.kind)
        }
        # Binding và policy đã được execution kiểm tra trước khi gọi capture.
        self.pass_rules("IN-COM-001", "IN-COM-002")
        self.issues: list[ValidationIssue] = []
        self.rows = 0
        self.blank_rows = 0
        self.scan_complete = False
        self.buffer_bytes = 0
        self.row_number: int | None = None
        self.column_number: int | None = None

    def pass_rules(self, *rules: str) -> None:
        for rule in rules:
            if self.checks[rule].status == "NOT_RUN":
                self.checks[rule] = RuleCheck(rule_id=rule, status="PASS")

    def warning(self, rule: str, code: ErrorCode) -> None:
        self.checks[rule] = RuleCheck(rule_id=rule, status="WARNING")
        self.issues.append(ValidationIssue(rule_id=rule, code=code, severity="WARNING"))

    def failure_report(self, exc: Exception) -> ValidationReport:
        if isinstance(exc, InputValidationError):
            violation = exc
        elif isinstance(exc, FileNotFoundError):
            violation = InputValidationError("IN-COM-003", "SOURCE_NOT_FOUND")
        elif isinstance(exc, PermissionError):
            violation = InputValidationError(
                "IN-COM-003", "SOURCE_ACCESS_DENIED", infrastructure=True
            )
        elif isinstance(exc, UnicodeError):
            violation = InputValidationError("IN-CSV-004", "ENCODING_ERROR")
        elif isinstance(exc, (csv.Error, BadZipFile)):
            violation = InputValidationError("IN-COM-004", "FILE_PARSE_ERROR")
        elif isinstance(exc, S3Error):
            missing = exc.code in {"NoSuchKey", "NoSuchBucket", "NoSuchVersion"}
            violation = InputValidationError(
                "IN-COM-003",
                "SOURCE_NOT_FOUND"
                if missing
                else (
                    "SOURCE_ACCESS_DENIED" if exc.code == "AccessDenied" else "SOURCE_UNAVAILABLE"
                ),
                infrastructure=not missing,
            )
        elif isinstance(exc, psycopg.Error):
            codes: dict[str, ErrorCode] = {
                "42P01": "SOURCE_NOT_FOUND",
                "42501": "SOURCE_ACCESS_DENIED",
                "28P01": "SOURCE_ACCESS_DENIED",
                "57014": "RESOURCE_LIMIT_EXCEEDED",
            }
            code = codes.get(exc.sqlstate or "", "SOURCE_UNAVAILABLE")
            violation = InputValidationError(
                "IN-COM-005" if exc.sqlstate == "57014" else "IN-PG-003",
                code,
                infrastructure=exc.sqlstate not in {"42P01", "57014"},
            )
        else:
            violation = InputValidationError(
                "IN-COM-004", "VALIDATION_EXECUTION_ERROR", infrastructure=True
            )
        status: Literal["ERROR", "FAIL"] = "ERROR" if violation.infrastructure else "FAIL"
        self.checks[violation.rule] = RuleCheck(rule_id=violation.rule, status=status)
        self.issues.append(
            ValidationIssue(
                rule_id=violation.rule,
                code=violation.code,
                severity="ERROR",
                row_number=self.row_number,
                column_number=self.column_number,
            )
        )
        return self.report()

    def report(self) -> ValidationReport:
        return ValidationReport(
            report_id=str(uuid4()),
            ingestion_job_id=self.job_id,
            attempt_id=f"job-{self.job_id}",
            policy_id=self.policy.policy_id,
            policy_version=self.policy.version,
            policy_fingerprint=self.policy.fingerprint,
            source_type=self.policy.parser.kind,
            scan_complete=self.scan_complete,
            records_scanned=self.rows,
            blank_rows_skipped=self.blank_rows,
            checks=tuple(self.checks.values()),
            issues=tuple(self.issues),
            total_errors=sum(i.severity == "ERROR" for i in self.issues),
            total_warnings=sum(i.severity == "WARNING" for i in self.issues),
        )

    def budget(self, *, size: int = 0) -> None:
        limits = self.policy.limits
        if (
            size > limits.max_file_bytes
            or self.rows > limits.max_records
            or self.buffer_bytes > limits.memory_limit_bytes
            or time.monotonic() - self.started > limits.timeout_seconds
        ):
            raise InputValidationError("IN-COM-005", "RESOURCE_LIMIT_EXCEEDED")

    def columns(self, names: list[str]) -> None:
        if not names or any(not name.strip() for name in names):
            raise InputValidationError("IN-COM-006", "HEADER_EMPTY")
        if len(set(names)) != len(names):
            raise InputValidationError("IN-COM-006", "HEADER_DUPLICATE")
        if len(names) > self.policy.limits.max_columns:
            raise InputValidationError("IN-COM-005", "RESOURCE_LIMIT_EXCEEDED")
        expected = [c.name for c in self.policy.columns]
        if any(c.required and c.name not in names for c in self.policy.columns):
            raise InputValidationError("IN-COM-006", "SOURCE_SCHEMA_MISMATCH")
        extra = set(names) - set(expected)
        if extra and not self.policy.allow_extra_columns:
            raise InputValidationError("IN-COM-006", "SOURCE_SCHEMA_MISMATCH")
        if self.policy.column_order == "strict" and (
            [n for n in names if n in expected] != [n for n in expected if n in names]
        ):
            raise InputValidationError("IN-COM-006", "SOURCE_SCHEMA_MISMATCH")
        if extra:
            self.warning("IN-COM-006", "EXTRA_COLUMNS_ALLOWED")
        self.pass_rules("IN-COM-006")

    def row(self, values: tuple[Any, ...] | list[Any], width: int) -> None:
        if self.policy.parser.kind != "POSTGRESQL":
            self.buffer_bytes = sys.getsizeof(values) + sum(sys.getsizeof(v) for v in values)
        self.budget()
        if len(values) != width:
            rule = {"CSV": "IN-CSV-005", "EXCEL": "IN-XLS-006", "POSTGRESQL": "IN-PG-004"}[
                self.policy.parser.kind
            ]
            raise InputValidationError(rule, "SOURCE_SCHEMA_MISMATCH")
        for index, value in enumerate(values, 1):
            self.column_number = index
            if isinstance(value, (str, bytes)) and len(value) > self.policy.limits.max_cell_chars:
                raise InputValidationError("IN-COM-005", "RESOURCE_LIMIT_EXCEEDED")
        self.column_number = None
        self.rows += 1
        self.budget()

    def finish_scan(self) -> None:
        self.row_number = self.column_number = None
        self.scan_complete = True
        if self.rows == 0:
            if not self.policy.allow_empty_records:
                raise InputValidationError("IN-COM-007", "EMPTY_SOURCE")
            self.warning("IN-COM-007", "EMPTY_SOURCE_ALLOWED")
        self.pass_rules("IN-COM-005", "IN-COM-007")

    def file(self, path: Path) -> None:
        self.budget(size=path.stat().st_size)
        self.pass_rules("IN-COM-003")
        if path.stat().st_size == 0:
            raise InputValidationError("IN-COM-004", "HEADER_EMPTY")
        parser = self.policy.parser
        if isinstance(parser, CSVOptions):
            if path.suffix.lower() != ".csv":
                raise InputValidationError("IN-COM-004", "UNSUPPORTED_FILE_TYPE")
            self.csv(path, parser)
        elif isinstance(parser, ExcelOptions):
            if path.suffix.lower() not in {".xlsx", ".xlsm"}:
                raise InputValidationError("IN-COM-004", "UNSUPPORTED_FILE_TYPE")
            self.excel(path, parser)
        else:
            raise InputValidationError("IN-COM-002", "SOURCE_SCHEMA_MISMATCH")
        self.pass_rules("IN-COM-004")
        self.finish_scan()

    def csv(self, path: Path, parser: CSVOptions) -> None:
        # Không để DuckDB tự đổi header trùng hoặc tự đệm cột thiếu.
        # newline='' giữ escaped newline; csv.reader đếm bản ghi logic.
        with path.open(encoding=parser.encoding, newline="") as stream:
            reader = csv.reader(
                stream, delimiter=parser.delimiter, quotechar=parser.quote_char, strict=True
            )
            for _ in range(parser.skip_rows):
                next(reader, None)
                self.budget()
            self.pass_rules("IN-CSV-001")
            names = next(reader, []) if parser.has_header else list(parser.positional_columns)
            self.columns(names)
            self.pass_rules("IN-CSV-002", "IN-CSV-003")
            for row in reader:
                self.row_number = self.rows + 1
                # Dòng trắng là sai cấu trúc; không âm thầm bỏ qua.
                self.row(row, len(names))
            self.pass_rules("IN-CSV-004", "IN-CSV-005")

    def excel(self, path: Path, parser: ExcelOptions) -> None:
        book = load_workbook(path, read_only=True, data_only=False)
        iterator = None
        try:
            self.pass_rules("IN-XLS-001", "IN-XLS-003")
            if parser.sheet_name not in book.sheetnames:
                raise InputValidationError("IN-XLS-002", "SHEET_NOT_FOUND")
            self.pass_rules("IN-XLS-002")
            sheet = book[parser.sheet_name]
            # Không tin dimension khai báo trong XML: có thể che mất dòng/cột thực tế.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            for _ in range(parser.skip_rows + parser.header_row - 1):
                next(iterator, None)
                self.budget()
            header = next(iterator, ())
            self.row_number = parser.skip_rows + parser.header_row
            self.excel_cells(header)
            names = [str(c.value).strip() if c.value is not None else "" for c in header]
            while names and not names[-1]:
                names.pop()
            self.columns(names)
            self.pass_rules("IN-XLS-004")
            for number, cells in enumerate(iterator, self.row_number + 1):
                self.row_number = number
                self.budget()
                self.excel_cells(cells)
                if any(c.value is not None for c in cells[len(names) :]):
                    raise InputValidationError("IN-XLS-006", "EXCEL_UNNAMED_COLUMN")
                values = [c.value for c in cells[: len(names)]]
                if all(v is None for v in values):
                    self.blank_rows += 1
                    continue
                values.extend([None] * (len(names) - len(values)))
                self.row(values, len(names))
            self.pass_rules("IN-XLS-005", "IN-XLS-006", "IN-XLS-007")
        finally:
            if iterator is not None:
                iterator.close()
            book.close()

    def excel_cells(self, cells: Any) -> None:
        for index, cell in enumerate(cells, 1):
            self.column_number = index
            if cell.data_type in {"f", "e"}:
                code: ErrorCode = (
                    "EXCEL_FORMULA_UNSUPPORTED" if cell.data_type == "f" else "EXCEL_CELL_ERROR"
                )
                raise InputValidationError("IN-XLS-005", code)
        self.column_number = None

    def postgres(self, request: ReadRequest, path: Path) -> None:
        parser = self.policy.parser
        assert isinstance(parser, PostgreSQLOptions)
        config = request.configuration
        # Kết nối không được lưu trong report. Transaction read-only REPEATABLE READ
        # giữ schema và các batch cùng snapshot; server cursor tránh fetchall.
        with psycopg.connect(
            host=config["host"],
            port=config["port"],
            dbname=config["database"],
            user=config["username"],
            password=config["password"],
            connect_timeout=self.policy.limits.timeout_seconds,
            options=f"-c statement_timeout={self.policy.limits.timeout_seconds * 1000}",
        ) as conn:
            conn.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
            self.pass_rules("IN-PG-001", "IN-PG-002", "IN-COM-003")
            query = sql.SQL("SELECT * FROM {}.{}").format(
                sql.Identifier(parser.schema_name), sql.Identifier(parser.table_name)
            )
            # LIMIT 0 lấy kiểu kể cả bảng rỗng, cùng transaction với server cursor.
            description = conn.execute(query + sql.SQL(" LIMIT 0")).description
            assert description is not None
            names = [c.name for c in description]
            self.columns(names)
            oids = [c.type_code for c in description]
            type_rows = conn.execute(
                "SELECT oid, typname FROM pg_catalog.pg_type WHERE oid = ANY(%s)", (oids,)
            ).fetchall()
            types: dict[int, str] = dict(type_rows)
            expected = {c.name: c for c in self.policy.columns}
            for column in description:
                if column.name in expected and types[column.type_code] not in (
                    expected[column.name].allowed_source_types
                ):
                    raise InputValidationError("IN-PG-004", "SOURCE_SCHEMA_MISMATCH")
            schema = pa.schema(
                [pa.field(c.name, postgres_arrow_type(types[c.type_code], c)) for c in description]
            )
            self.pass_rules("IN-PG-003", "IN-PG-004", "IN-COM-004")
            with (
                conn.cursor(name="bronze_validation") as cursor,
                pq.ParquetWriter(path, schema) as writer,
            ):
                cursor.itersize = 512
                cursor.execute(query)
                while True:
                    batch: list[tuple[Any, ...]] = []
                    self.buffer_bytes = 0
                    for values in islice(cursor, 512):
                        self.row_number = self.rows + 1
                        self.buffer_bytes += sys.getsizeof(values) + sum(
                            sys.getsizeof(v) for v in values
                        )
                        self.row(values, len(names))
                        batch.append(values)
                    if not batch:
                        break
                    writer.write_table(
                        pa.Table.from_arrays(
                            [
                                pa.array([r[i] for r in batch], type=f.type)
                                for i, f in enumerate(schema)
                            ],
                            schema=schema,
                        )
                    )
                    self.budget(size=path.stat().st_size)
            self.budget(size=path.stat().st_size)
            self.pass_rules("IN-PG-005", "IN-PG-006", "IN-COM-008")
            self.finish_scan()


def postgres_arrow_type(name: str, column: Any) -> Any:
    """Tập kiểu được hỗ trợ tường minh; không suy kiểu bằng mẫu hoặc biến kiểu lạ thành text."""
    types: dict[str, Any] = {
        "int2": pa.int16(),
        "int4": pa.int32(),
        "int8": pa.int64(),
        "bool": pa.bool_(),
        "float4": pa.float32(),
        "float8": pa.float64(),
        "text": pa.string(),
        "varchar": pa.string(),
        "bpchar": pa.string(),
        "bytea": pa.binary(),
        "date": pa.date32(),
        "timestamp": pa.timestamp("us"),
        "timestamptz": pa.timestamp("us", tz="UTC"),
        "time": pa.time64("us"),
    }
    if (
        name == "numeric"
        and column.precision
        and column.scale is not None
        and 0 <= column.scale <= column.precision <= 76
    ):
        return (pa.decimal128 if column.precision <= 38 else pa.decimal256)(
            column.precision, column.scale
        )
    if name in types:
        return types[name]
    raise InputValidationError("IN-PG-004", "SOURCE_SCHEMA_MISMATCH")
