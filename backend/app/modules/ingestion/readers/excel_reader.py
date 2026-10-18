from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .contract import ReadRequest, ReadResult

SUPPORTED_EXCEL_SUFFIXES = {".xlsx", ".xlsm"}


class ExcelReader:
    """Đọc dữ liệu Excel bằng openpyxl."""

    def read(self, request: ReadRequest) -> ReadResult:
        """Đọc một sheet trong file Excel local và chuẩn hóa về ReadResult."""

        # 1. Reader này chỉ xử lý nguồn Excel.
        # Dùng upper() để caller truyền "excel" hay "EXCEL" đều đọc được.
        if request.source_type.upper() != "EXCEL":
            raise ValueError("Source type phải là EXCEL")

        # 2. location là đường dẫn file Excel hoặc thư mục chứa file Excel.
        if not request.location:
            raise ValueError("Location không được để trống")

        # 3. Lấy cấu hình parser từ data_source.configuration.
        # sheet_name: sheet cần đọc; header_row: dòng header; skip_rows: số dòng metadata bỏ qua.
        config = request.configuration

        sheet_name = _required_text(config, "sheet_name")
        header_row = _required_int(config, "header_row", minimum=1)
        skip_rows = _required_int(config, "skip_rows", minimum=0)

        # 4. Xác định đường dẫn file cần đọc.
        # Reader này chỉ đọc file local; S3/HTTP sẽ cần reader riêng.
        if "://" in request.location:
            raise ValueError("ExcelReader hiện chỉ hỗ trợ file local")

        path = Path(request.location).expanduser()

        if path.is_dir():
            if not request.source_table:
                raise ValueError("Thiếu source_table")

            # Khi location là thư mục, source_table phải là tên file trực tiếp.
            # Không nhận đường dẫn con để tránh đọc nhầm file ngoài thư mục nguồn.
            if Path(request.source_table).name != request.source_table:
                raise ValueError("source_table không hợp lệ")

            path = path / request.source_table

        elif request.source_table and path.name != request.source_table:
            raise ValueError("source_table không khớp với location")

        if not path.is_file():
            raise FileNotFoundError(f"Không tìm thấy file Excel: {path}")

        if path.suffix.lower() not in SUPPORTED_EXCEL_SUFFIXES:
            raise ValueError("Chỉ hỗ trợ file .xlsx và .xlsm")

        # Đọc công thức tường minh: kết quả lưu đệm có thể thiếu hoặc đã cũ.
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
        )

        iterator = None
        try:
            # 6. Kiểm tra sheet trước khi đọc dữ liệu.
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Không tìm thấy sheet: {sheet_name}")

            sheet = workbook[sheet_name]
            iterator = sheet.iter_rows()

            # header_row được tính sau skip_rows.
            # Ví dụ skip_rows=2, header_row=1 nghĩa là bỏ 2 dòng metadata,
            # rồi lấy dòng tiếp theo làm header.
            rows_to_skip = skip_rows + header_row - 1

            for _ in range(rows_to_skip):
                next(iterator, None)

            # 7. Đọc và chuẩn hóa header.
            header_values = next(iterator, None)

            if header_values is None:
                raise ValueError("Không tìm thấy dòng header")

            if any(cell.data_type in {"f", "e"} for cell in header_values):
                raise ValueError("EXCEL_UNSUPPORTED_HEADER_CELL")
            columns = _normalize_columns(tuple(cell.value for cell in header_values))

            if not columns or any(not col for col in columns):
                raise ValueError("Header chứa tên cột trống")

            if len(columns) != len(set(columns)):
                raise ValueError("Header chứa tên cột trùng nhau")

            # 8. Đọc các dòng dữ liệu theo đúng số cột header.
            # Bỏ dòng hoàn toàn trống để tránh sinh record rỗng ở cuối sheet.
            rows = []
            locators = []
            for row_number, row in enumerate(iterator, start=rows_to_skip + 2):
                # Đánh dấu batch lỗi trước khi ánh xạ, kể cả các dòng chỉ chứa công thức.
                # Không đưa giá trị thô hoặc nội dung công thức vào thông báo lỗi.
                for cell in row:
                    if cell.data_type in {"f", "e"}:
                        code = (
                            "EXCEL_FORMULA_UNSUPPORTED"
                            if cell.data_type == "f"
                            else "EXCEL_CELL_ERROR"
                        )
                        raise ValueError(f"{code}: row={row_number}, cell={cell.coordinate}")
                if any(cell.value is not None for cell in row[len(columns) :]):
                    raise ValueError(f"EXCEL_UNNAMED_COLUMN: row={row_number}")
                normalized_row = tuple(cell.value for cell in row[: len(columns)])

                if all(value is None for value in normalized_row):
                    continue

                rows.append(normalized_row)
                locators.append(f"sheet:{sheet_name}:row:{row_number}")

            # 9. Trả dữ liệu về format chung của ingestion reader.
            return ReadResult(
                columns=columns,
                rows=rows,
                source_table=request.source_table or path.name,
                record_locators=tuple(locators),
                workbook_epoch="1904" if workbook.epoch.year == 1904 else "1900",
            )

        finally:
            if iterator is not None:
                iterator.close()
            workbook.close()


def _required_text(config: dict[str, Any], key: str) -> str:
    """Lấy một field cấu hình dạng chuỗi và báo lỗi nếu thiếu hoặc rỗng."""

    value = config.get(key)

    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{key} không hợp lệ")

    return value.strip()


def _required_int(config: dict[str, Any], key: str, *, minimum: int) -> int:
    """Lấy một field cấu hình dạng int và kiểm tra giá trị nhỏ nhất."""

    value = config.get(key)

    # bool là subclass của int trong Python, nên cần loại riêng True/False.
    if isinstance(value, bool) or not isinstance(value, int) or value < minimum:
        raise ValueError(f"{key} phải >= {minimum}")

    return value


def _normalize_columns(values: tuple[Any, ...]) -> list[str]:
    """Chuyển dòng header trong Excel thành danh sách tên cột sạch."""

    columns = [str(value).strip() if value is not None else "" for value in values]

    # Bỏ các ô header trống ở cuối vì chúng thường chỉ là vùng Excel dư.
    while columns and not columns[-1]:
        columns.pop()

    return columns
