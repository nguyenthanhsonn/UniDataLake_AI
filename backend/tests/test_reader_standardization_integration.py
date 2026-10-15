"""Real local files through reader -> Bronze hand-off -> S2 -> S3."""

from __future__ import annotations

import csv
import hashlib
import json
from dataclasses import replace
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from app.modules.ingestion.contracts import MappingBinding
from app.modules.ingestion.reader_batch import reader_to_batch
from app.modules.ingestion.readers.contract import ReadRequest, ReadResult
from app.modules.ingestion.readers.csv_reader import CSVReader
from app.modules.ingestion.readers.excel_reader import ExcelReader
from app.modules.ingestion.readers.file_snapshot import read_file_snapshot
from app.modules.pipeline.identity_contracts import CrosswalkEntry, CrosswalkSnapshot
from app.modules.pipeline.identity_standardization import (
    IdentityStandardizationService,
    ValueDictionary,
)
from app.modules.pipeline.mapping_profiles import build_profile, pointer
from app.modules.pipeline.schema_mapping import MappingKey, MappingRegistry, SchemaMappingService
from app.modules.pipeline.type_standardization import (
    ConversionConfigurationError,
    FieldOptions,
    TypeStandardizationService,
)

ROOT = Path(__file__).resolve().parents[2]
RECORDS = json.loads(
    (ROOT / "data/standardization/s1/expected-business-records.json").read_text(encoding="utf-8")
)


@pytest.mark.parametrize("transport", ["CSV", "EXCEL"])
@pytest.mark.parametrize(
    ("domain", "entity", "pk", "raw_id", "canonical"),
    [
        ("admissions", "admission_method", "admission_method_id", "000203", 9001),
        ("academic", "semester", "semester_id", "000305", 9002),
        ("hr", "position", "position_id", "000102", 9003),
    ],
)
def test_reader_s2_s3_s4_preserves_provenance_and_resolves_ids(
    tmp_path, transport, domain, entity, pk, raw_id, canonical
):
    read, mapped, typed = run_file(tmp_path, domain, entity, transport, {pk: raw_id})
    contract = json.loads(
        (ROOT / f"docs/standardization/silver-schemas/{domain}.json").read_text(encoding="utf-8")
    )
    types = TypeStandardizationService(
        key=mapped.key, conversion_version="integration-1", contract=contract
    )
    when = datetime(2026, 1, 1, tzinfo=UTC)
    engine = IdentityStandardizationService(
        types=types,
        version="fixture-s4-1",
        crosswalk=CrosswalkSnapshot(
            "fixture-cw-1",
            (
                CrosswalkEntry(
                    mapped.key.source_system_id, entity, (raw_id,), canonical, "fixture-only", when
                ),
            ),
        ),
        dictionaries={
            n: ValueDictionary("fixture-dict-1")
            for n, spec in types.spec["properties"].items()
            if "enum" in spec
        },
        reference_namespaces={},
        effective_at=when,
    )
    result = engine.standardize(typed)
    assert not result.issues
    assert result.accepted[0].values[pk] == canonical
    assert result.accepted[0].locator == typed.accepted[0].locator
    assert result.provenance == read.provenance
    assert read.provenance is not None
    assert result.input_version == read.provenance.input_version
    assert not result.silver_ready


def run_file(tmp_path, domain, entity, transport="EXCEL", changes=None, options=None, epoch=None):
    record = {**RECORDS[entity][0], **(changes or {})}
    columns = ["source/" + name for name in record]
    path = tmp_path / ("input.xlsx" if transport == "EXCEL" else "input.csv")
    if transport == "EXCEL":
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        if epoch:
            workbook.epoch = epoch
        sheet.append(["metadata"])
        sheet.append(columns)
        sheet.append([None] * len(columns))
        sheet.append(list(record.values()))
        workbook.save(path)
        workbook.close()
        config = {"sheet_name": "Data", "header_row": 1, "skip_rows": 1}
        reader = ExcelReader()
    else:
        with path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(columns)
            writer.writerow([str(v).lower() if isinstance(v, bool) else v for v in record.values()])
        config = {"delimiter": ",", "encoding": "utf-8", "has_header": True}
        reader = CSVReader()
    version = "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()
    read = read_file_snapshot(
        ReadRequest(transport, str(path), config),
        reader,
        source_system_id="integration-source",
        data_source_id="1",
        dataset_id=entity,
    )
    contract = json.loads(
        (ROOT / f"docs/standardization/silver-schemas/{domain}.json").read_text(encoding="utf-8")
    )
    key = MappingKey(
        "integration-source", domain, entity, "1", "reader-map", "1", contract["x-version"]
    )
    batch = reader_to_batch(
        read,
        batch_id="integration",
        domain_id=domain,
        source_type=transport,
        input_version=version,
        binding=MappingBinding(
            key.source_system_id, entity, "1", key.mapping_id, "1", key.target_schema_version
        ),
    )
    target = contract["$defs"][entity]["properties"]
    profile = build_profile(
        key=key,
        source_type=transport,
        source_columns=batch.source_paths,
        target_fields=tuple(target),
        field_paths={
            name: (pointer(column),) for name, column in zip(record, columns, strict=True)
        },
        deferred_fields=tuple(name for name in target if name not in record),
    )
    mapped = SchemaMappingService(MappingRegistry((profile,))).map_batch(batch)
    typed = TypeStandardizationService(
        key=key, conversion_version="integration-1", contract=contract, options=options
    ).standardize(mapped)
    assert typed.input_version == version
    assert typed.provenance == read.provenance
    assert typed.provenance is not None
    assert typed.provenance.dataset_id == entity
    assert not typed.silver_ready
    assert typed.total_count == 1
    assert len(typed.accepted) + len(typed.rejected_locators) == 1
    assert mapped.accepted[0].locator == (
        "sheet:Data:row:4" if transport == "EXCEL" else "record:1"
    )
    return read, mapped, typed


@pytest.mark.parametrize("transport", ["CSV", "EXCEL"])
@pytest.mark.parametrize(
    ("domain", "entity", "field", "value", "expected"),
    [
        ("admissions", "admission_method", "max_score", "29.50", Decimal("29.50")),
        ("academic", "semester", "start_date", "2026-09-01", date(2026, 9, 1)),
        ("hr", "position", "position_id", "0007", "0007"),
    ],
)
def test_three_domains_real_files(tmp_path, transport, domain, entity, field, value, expected):
    _, mapped, typed = run_file(tmp_path, domain, entity, transport, {field: value})
    assert not mapped.issues
    assert not typed.issues
    row = typed.accepted[0]
    actual = row.pending_s4[field] if field in row.pending_s4 else row.values[field]
    assert actual == expected
    assert type(actual) is type(expected)
    assert typed.accepted[0].pending_s4


@pytest.mark.parametrize(
    ("domain", "entity", "field", "value", "code"),
    [
        ("admissions", "admission_method", "max_score", 29.5, "DECIMAL_PARSE_ERROR"),
        ("academic", "semester", "start_date", date(2026, 9, 1), "DATE_TIME_LOSS"),
        ("academic", "semester", "start_date", datetime(2026, 9, 1, 12, 30), "DATE_TIME_LOSS"),
        ("hr", "position", "position_code", 7, "STRING_TYPE_ERROR"),
        ("hr", "position", "position_code", None, "REQUIRED_VALUE_MISSING"),
        ("academic", "semester", "start_date", "31/02/2026", "DATE_PARSE_ERROR"),
    ],
)
def test_excel_unsafe_values_rejected_with_locator(tmp_path, domain, entity, field, value, code):
    _, _, typed = run_file(tmp_path, domain, entity, changes={field: value})
    assert not typed.accepted
    assert typed.rejected_locators == ("sheet:Data:row:4",)
    assert [(i.field, i.code, i.locator) for i in typed.issues] == [
        (field, code, "sheet:Data:row:4")
    ]


def test_excel_locale_boolean_and_decimal_text(tmp_path):
    _, _, typed = run_file(
        tmp_path,
        "admissions",
        "admission_method",
        changes={"max_score": "29,50", "is_active": "Có"},
        options={
            "max_score": FieldOptions(decimal_separator=","),
            "is_active": FieldOptions(true_tokens=("Có",), false_tokens=("Không",)),
        },
    )
    assert typed.accepted[0].values["max_score"] == Decimal("29.50")
    assert typed.accepted[0].values["is_active"] is True


def test_s3_rejects_version_changed_after_s2(tmp_path):
    _, mapped, _ = run_file(tmp_path, "hr", "position")
    contract = json.loads(
        (ROOT / "docs/standardization/silver-schemas/hr.json").read_text(encoding="utf-8")
    )
    service = TypeStandardizationService(
        key=mapped.key, conversion_version="integration-1", contract=contract
    )
    with pytest.raises(ConversionConfigurationError, match="PROVENANCE_MISMATCH"):
        service.standardize(replace(mapped, input_version="wrong"))


@pytest.mark.parametrize(
    ("serial", "epoch", "label", "expected"),
    [
        (61, None, "1900", date(1900, 3, 1)),
        (0, CALENDAR_MAC_1904, "1904", date(1904, 1, 1)),
    ],
)
def test_excel_serial_date_explicit_epoch(tmp_path, serial, epoch, label, expected):
    read, _, typed = run_file(
        tmp_path,
        "academic",
        "semester",
        changes={"start_date": serial},
        epoch=epoch,
        options={"start_date": FieldOptions(excel_epoch=label)},
    )
    assert typed.accepted[0].values["start_date"] == expected
    assert read.provenance is not None
    assert read.provenance.workbook_epoch == label


@pytest.mark.parametrize(
    ("value", "code"), [("=1+1", "EXCEL_FORMULA_UNSUPPORTED"), ("#DIV/0!", "EXCEL_CELL_ERROR")]
)
def test_excel_formula_or_error_never_becomes_null(tmp_path, value, code):
    with pytest.raises(ValueError, match=code):
        run_file(tmp_path, "admissions", "admission_method", changes={"max_score": value})


def test_excel_native_timestamp_with_explicit_timezone(tmp_path):
    _, _, typed = run_file(
        tmp_path,
        "hr",
        "department",
        changes={"created_at": datetime(2026, 1, 1, 7)},
        options={"created_at": FieldOptions(source_timezone="Asia/Ho_Chi_Minh")},
    )
    assert not typed.issues
    assert typed.accepted[0].values["created_at"] == datetime(2026, 1, 1, tzinfo=UTC)


def test_excel_nullable_blank_is_preserved(tmp_path):
    _, _, typed = run_file(tmp_path, "hr", "assignment")
    assert not typed.issues
    assert typed.accepted[0].values["end_date"] is None


def test_excel_data_without_header_is_not_silently_truncated(tmp_path):
    path = tmp_path / "extra.xlsx"
    workbook = Workbook()
    assert workbook.active is not None
    workbook.active.append(["value"])
    workbook.active.append([1, "unmapped"])
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="EXCEL_UNNAMED_COLUMN: row=2"):
        ExcelReader().read(
            ReadRequest(
                "EXCEL", str(path), {"sheet_name": "Sheet", "header_row": 1, "skip_rows": 0}
            )
        )


def test_excel_formula_only_row_is_not_dropped(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    assert workbook.active is not None
    workbook.active.append(["value"])
    workbook.active.append(["=1+1"])
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="EXCEL_FORMULA_UNSUPPORTED: row=2"):
        ExcelReader().read(
            ReadRequest(
                "EXCEL", str(path), {"sheet_name": "Sheet", "header_row": 1, "skip_rows": 0}
            )
        )


@pytest.mark.parametrize(
    ("columns", "rows", "code"),
    [
        (["a", "a"], [(1, 2)], "DUPLICATE_READER_COLUMNS"),
        (["a"], [(1, 2)], "READER_ROW_WIDTH_MISMATCH"),
        ([""], [(1,)], "INVALID_READER_COLUMNS"),
    ],
)
def test_adapter_rejects_lossy_shapes(columns, rows, code):
    with pytest.raises(ValueError, match=code):
        reader_to_batch(
            ReadResult(columns, rows, None),
            batch_id="b",
            domain_id="hr",
            source_type="EXCEL",
            input_version="v1",
            binding=MappingBinding("s", "position", "1", "m", "1", "1"),
        )
